"""Independently validate real PDF, Word and image export containers."""
from pathlib import Path
import io
import json
import sys
import zipfile
from xml.etree import ElementTree

from docx import Document
from PIL import Image
from pypdf import PdfReader
from openpyxl import load_workbook


def main(directory):
    results = json.loads((directory / "results.json").read_text(encoding="utf8"))
    checks = []
    for filename in ["semidata-report-pdf.pdf", "production-report.pdf", "production-all-sections.pdf"]:
        reader = PdfReader(directory / filename, strict=True)
        assert len(reader.pages) >= 1
        assert all(len(page.images) == 1 for page in reader.pages)
        receipt = json.loads(reader.attachments["report-receipt.json"][0])
        assert receipt["format"] == "semidata-document-report"
        assert receipt["provenance"]["sources"]
        if filename.startswith("semidata"):
            assert len(reader.pages) == results["files"][0]["pages"]
            assert receipt["provenance"]["sources"][0]["name"] == "시험 测试.stdf"
        checks.append({"file": filename, "pages": len(reader.pages), "receipt": True})
    for filename in ["semidata-report-docx.docx", "production-report.docx"]:
        document = Document(directory / filename)
        assert document.paragraphs[0].style.name == "Title"
        assert len(document.tables) >= 1
        assert len(document.inline_shapes) >= 1
        with zipfile.ZipFile(directory / filename) as package:
            assert package.testzip() is None
            for name in package.namelist():
                if name.endswith((".xml", ".rels")):
                    ElementTree.fromstring(package.read(name))
            receipt = json.loads(package.read("report-receipt.json"))
            if filename.startswith("semidata"):
                assert any("측정 결과 测试结果" in paragraph.text for paragraph in document.paragraphs)
                assert len(document.tables[1].rows) == 53
                assert [cell.text for cell in document.tables[0].rows[0].cells] == ["Measurement", "Value", "Unit"]
                assert [cell.text for cell in document.tables[0].rows[1].cells] == ["Mean", "1.25", "μA"]
        checks.append({"file": filename, "tables": len(document.tables), "images": len(document.inline_shapes), "receipt": True})
    for extension in ["png", "jpg"]:
        filename = f"semidata-report-{extension}.zip"
        with zipfile.ZipFile(directory / filename) as package:
            images = [name for name in package.namelist() if name.endswith(f".{extension}")]
            assert len(images) == results["files"][0]["pages"]
            for name in images:
                image = Image.open(io.BytesIO(package.read(name)))
                assert image.size == (1200, 1697)
                assert image.format == ("PNG" if extension == "png" else "JPEG")
                image.verify()
            assert "report-receipt.json" in package.namelist()
        checks.append({"file": filename, "images": len(images)})
    landscape_pdf = PdfReader(directory / "semidata-landscape-pdf.pdf", strict=True)
    assert len(landscape_pdf.pages) == 1
    assert abs(float(landscape_pdf.pages[0].mediabox.width) - 297 * 72 / 25.4) < 0.001
    assert abs(float(landscape_pdf.pages[0].mediabox.height) - 210 * 72 / 25.4) < 0.001
    landscape_docx = Document(directory / "semidata-landscape-docx.docx")
    assert abs(landscape_docx.sections[0].page_width.mm - 297) < 0.02
    assert abs(landscape_docx.sections[0].page_height.mm - 210) < 0.02
    landscape_png = Image.open(directory / "semidata-landscape-png.png")
    assert landscape_png.size == (1600, 1131)
    receipt = json.loads(landscape_pdf.attachments["report-receipt.json"][0])
    assert receipt["layout"] == {"pageWidthMm": 297, "pageHeightMm": 210, "imageWidthPx": 1600}
    checks.append({"layout": "297 × 210 mm", "PDF/DOCX geometry": True, "image": [1600, 1131]})
    for filename in ["semidata-summary-xlsx.xlsx", "production-report.xlsx", "batch-lot-report.xlsx"]:
        workbook = load_workbook(directory / filename)
        assert "Provenance" in workbook.sheetnames
        assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
        rows = list(workbook["Provenance"].values)
        receipt_text = "".join(row[1] for row in rows if str(row[0]).startswith("Receipt JSON "))
        receipt = json.loads(receipt_text)
        assert receipt["format"] == "semidata-document-report"
        sources = [row for row in rows if row[0] == "Source"]
        if filename.startswith("semidata"):
            literals = workbook["Literal cells"]
            assert literals["A4"].value == '=HYPERLINK("https://invalid.example")'
            assert literals["A4"].data_type == "s" and literals["B4"].value == 12.5
            assert literals["A5"].value == "+SUM(1,2)" and literals["B5"].value is False
            assert literals["B6"].value == 0
            assert sum(len(sheet._images) for sheet in workbook) == 1
            assert len(sources) == 1
        elif filename.startswith("batch"):
            assert len(sources) == 2 and len(receipt["provenance"]["sources"]) == 2
            assert workbook["Selected population"]["B5"].value == 2
        else:
            assert sum(len(sheet._images) for sheet in workbook) == 10
            assert workbook["Selected population"]["B5"].value == 6
        with zipfile.ZipFile(directory / filename) as package:
            for path in package.namelist():
                if path.startswith("xl/media/"):
                    image = Image.open(io.BytesIO(package.read(path)))
                    assert image.format == "PNG"
                    image.verify()
        checks.append({"file": filename, "sheets": len(workbook.sheetnames), "sourceCount": len(sources), "literalCells": True})
    (directory / "container-validation.json").write_text(json.dumps(checks, indent=2), encoding="utf8")
    print(json.dumps(checks, indent=2))


if __name__ == "__main__":
    main(Path(sys.argv[1]))
